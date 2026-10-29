#!/usr/bin/env python3
"""Genera DOCX de manual a partir de manifiesto JSON y capturas anotadas."""
import argparse, json, os, shutil, subprocess
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

GREEN=RGBColor(0,96,70)
DARK=RGBColor(27,43,58)
GREY=RGBColor(85,101,118)
ORANGE=RGBColor(167,122,72)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd');shd.set(qn('w:fill'), fill);tcPr.append(shd)

def set_cell_margin(cell, top=120, start=120, bottom=100, end=120):
    tc=cell._tc;tcPr=tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
    for k,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        el=OxmlElement('w:'+k);el.set(qn('w:w'),str(v));el.set(qn('w:type'),'dxa');m.append(el)
    tcPr.append(m)

def add_field(paragraph, field):
    r=paragraph.add_run();beg=OxmlElement('w:fldChar');beg.set(qn('w:fldCharType'),'begin');inst=OxmlElement('w:instrText');inst.set(qn('xml:space'),'preserve');inst.text=' '+field+' ';sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');val=OxmlElement('w:t');val.text='1';end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');[r._r.append(x) for x in [beg,inst,sep,val,end]]

def add_label(cell, lab, text):
    p=cell.add_paragraph();p.paragraph_format.space_after=Pt(4)
    r=p.add_run(lab.upper());r.bold=True;r.font.size=Pt(8);r.font.color.rgb=ORANGE
    p=cell.add_paragraph(text);p.style='Normal';p.paragraph_format.space_after=Pt(12)
    for r in p.runs:r.font.size=Pt(10);r.font.color.rgb=DARK

def generate(path, out):
    path=Path(path).resolve();data=json.loads(path.read_text(encoding='utf-8'));out=Path(out).resolve();out.mkdir(parents=True,exist_ok=True)
    doc=Document();sec=doc.sections[0];sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Inches(11.69);sec.page_height=Inches(8.27);sec.top_margin=Inches(.67);sec.bottom_margin=Inches(.68);sec.left_margin=Inches(.71);sec.right_margin=Inches(.71)
    styles=doc.styles;normal=styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(10);normal.font.color.rgb=DARK;normal.paragraph_format.space_after=Pt(7)
    for sty,size,color in [('Title',31,GREEN),('Heading 1',21,GREEN),('Heading 2',14,GREEN)]:
        s=styles[sty];s.font.name='Aptos Display';s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=color;s.paragraph_format.space_after=Pt(10)
    h=sec.header.paragraphs[0];h.add_run('EcoLact  /  GESTIÓN INTEGRAL').bold=True;h.runs[0].font.size=Pt(9);h.runs[0].font.color.rgb=GREEN;h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=f.add_run('MANUAL DE USUARIO  •  DEMOSTRACIÓN VISUAL     |     ');r.font.size=Pt(8);r.font.color.rgb=GREY;add_field(f,'PAGE')
    # Cover
    p=doc.add_paragraph();p.space_before=Pt(38);r=p.add_run('MANUAL DE USUARIO');r.bold=True;r.font.size=Pt(14);r.font.color.rgb=ORANGE
    doc.add_paragraph(data['titulo'],style='Title')
    p=doc.add_paragraph(data['subtitulo']);p.style='Heading 2';p.paragraph_format.space_after=Pt(18)
    t=doc.add_table(rows=1,cols=2);t.autofit=False;t.columns[0].width=Inches(6.8);t.columns[1].width=Inches(3.3)
    t.cell(0,0).width=Inches(6.8);t.cell(0,1).width=Inches(3.3)
    first=path.parent/data['pasos'][1]['imagen_anotada'];t.cell(0,0).paragraphs[0].add_run().add_picture(str(first),width=Inches(6.6));t.cell(0,0).vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(t.cell(0,1),'F2F6F4');set_cell_margin(t.cell(0,1),210,240,180,230)
    add_label(t.cell(0,1),'Procedimiento',data['proceso']);add_label(t.cell(0,1),'Perfil',data['rol']);add_label(t.cell(0,1),'Versión',data['version']);add_label(t.cell(0,1),'Naturaleza','Ejemplo realizado con capturas proporcionadas. La creación final todavía no está verificada.')
    doc.add_page_break()
    # Scope and steps table
    doc.add_paragraph('Ficha del procedimiento',style='Heading 1')
    doc.add_paragraph('Propósito: orientar al administrador para abrir y completar el formulario de nueva zona de acopio. Se documenta lo visible sin asumir resultados no mostrados.')
    doc.add_paragraph('Antes de comenzar',style='Heading 2')
    for item in data.get('prerrequisitos',[]): doc.add_paragraph(item,style='List Bullet')
    doc.add_paragraph('Secuencia visual',style='Heading 2')
    table=doc.add_table(rows=1,cols=3);table.style='Light Shading Accent 1'
    head=table.rows[0].cells
    for c,v in zip(head,['Nº','Acción','Resultado / evidencia']):c.text=v
    for st in data['pasos']:
        cells=table.add_row().cells;cells[0].text=str(st['numero']);cells[1].text=st['titulo'];cells[2].text=st['estado']
    doc.add_paragraph('Criterio de validación',style='Heading 2')
    doc.add_paragraph('Para dar este procedimiento por concluido, se requiere una captura adicional después de confirmar «Crear zona» y comprobar la presencia del nuevo registro.')
    # Step-specific pages
    for st in data['pasos']:
        doc.add_page_break();p=doc.add_paragraph();p.style='Heading 1';p.add_run(f"Paso {st['numero']:02d}  |  {st['titulo']}")
        table=doc.add_table(rows=1,cols=2);table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
        table.columns[0].width=Inches(7.78);table.columns[1].width=Inches(2.22)
        c0,c1=table.rows[0].cells;c0.width=Inches(7.78);c1.width=Inches(2.22)
        img=path.parent/st['imagen_anotada']
        with Image.open(img) as im: width,height=im.size
        maxheight=5.85;natural_h=7.75*height/width;pic_width=min(7.75,maxheight*width/height)
        p0=c0.paragraphs[0];p0.add_run().add_picture(str(img),width=Inches(pic_width));p0.alignment=WD_ALIGN_PARAGRAPH.CENTER
        c0.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(c1,'F2F6F4');set_cell_margin(c1,170,180,180,150)
        add_label(c1,'Acción',st['instruccion'])
        add_label(c1,'Resultado',st['resultado'])
        add_label(c1,'Estado de evidencia',st['estado'])
        p=doc.add_paragraph(f"Figura {st['numero']}. Anotación sobre una captura real de EcoLact; el área marcada corresponde a la acción indicada.")
        p.style='Caption';p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()
    doc.add_paragraph('Evidencia pendiente y cierre',style='Heading 1')
    doc.add_paragraph('Las capturas actuales permiten identificar el botón y el formulario. No permiten afirmar que la zona haya sido guardada ni que exista una validación particular.')
    for x in data.get('pendientes',[]):doc.add_paragraph(x,style='List Bullet')
    doc.add_paragraph('Reglas de calidad para la documentación futura',style='Heading 2')
    for x in ['Captura original conservada junto a cada imagen anotada.','Solo una acción destacada por captura, con flecha legible.','Identidad del rol, precondición y resultado observables.','No presentar datos inventados como registros reales.','Verificación del guardado, estado y permisos antes de cerrar el procedimiento.']:
        doc.add_paragraph(x,style='List Bullet')
    doc.add_paragraph('Uso previsto: este ejemplo puede ampliarse a acopiadores, proveedores, entregas, calidad, liquidaciones y pagos una vez inspeccionado el sistema real.')
    outfile=out/'Manual_EcoLact_Zonas_Demostracion.docx';doc.save(outfile)
    print('DOCX:',outfile)
    office=shutil.which('libreoffice') or shutil.which('soffice')
    if not office and os.name=='nt':
        candidates=[Path(os.environ.get('ProgramFiles','C:/Program Files'))/'LibreOffice/program/soffice.exe',Path(os.environ.get('ProgramFiles(x86)','C:/Program Files (x86)'))/'LibreOffice/program/soffice.exe']
        office=next((str(q) for q in candidates if q.exists()),None)
    if office:
        cmd=[office,'-env:UserInstallation=file:///tmp/lo_manual_ecolact' if os.name!='nt' else '-env:UserInstallation=file:///C:/Users/Public/lo_manual_ecolact','--headless','--convert-to','pdf','--outdir',str(out),str(outfile)]
        try: subprocess.run(cmd,check=True,timeout=90);print('PDF:',outfile.with_suffix('.pdf'))
        except Exception as exc: print('PDF no convertido automáticamente:',exc)
    else: print('PDF pendiente: instala LibreOffice o exporta desde Microsoft Word.')

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('manifest');parser.add_argument('--out',default='salida');args=parser.parse_args();generate(args.manifest,args.out)
