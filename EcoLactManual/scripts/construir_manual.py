#!/usr/bin/env python3
"""Construye el manual completo del Administrador (Word + PDF) a partir de manual/manual.json.

El manual es ACUMULATIVO: manual.json lista todos los capítulos (un JSON por capítulo en
manual/capitulos/). Cada ejecución regenera el documento entero con todos ellos, así un capítulo
nuevo nunca reemplaza a los anteriores.

Uso: python scripts/construir_manual.py [--sin-pdf]
"""
import argparse
import json
import subprocess
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parent))
from anotacion import anotar  # noqa: E402

BASE = Path(__file__).resolve().parents[1]
MANUAL = BASE / 'manual'
ANOTADAS = BASE / 'evidencias' / 'anotadas'
SALIDA = BASE / 'salida'
NOMBRE = 'Manual_Usuario_EcoLact_Administrador'

# Paleta de la interfaz de EcoLactWeb (resources/css/app.css).
PRIMARIO = RGBColor(0x0F, 0x7C, 0x86)
PRIMARIO_OSC = RGBColor(0x0B, 0x59, 0x61)
ACENTO = RGBColor(0xA7, 0x7A, 0x48)
TEXTO = RGBColor(0x1F, 0x2D, 0x33)
TENUE = RGBColor(0x68, 0x77, 0x7D)
HEX_PRIMARIO_SUAVE = 'E1F2F1'
HEX_ACENTO_SUAVE = 'F2E7D8'
HEX_FONDO = 'F7F4EE'
HEX_EXITO_SUAVE = 'E8F8EF'
HEX_AVISO_SUAVE = 'FFF4DF'
FUENTE = 'Calibri'
ANCHO_UTIL = Cm(16.0)


# --------------------------------------------------------------------------- utilidades XML
def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def bordes_tabla(tabla, color='D5DDDF', tam=4, internos=True):
    tblPr = tabla._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    lados = ['top', 'left', 'bottom', 'right'] + (['insideH', 'insideV'] if internos else [])
    for lado in lados:
        e = OxmlElement(f'w:{lado}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(tam))
        e.set(qn('w:color'), color)
        b.append(e)
    tblPr.append(b)


def margenes_celda(celda, sup=90, izq=140, inf=90, der=140):
    tcPr = celda._tc.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for k, v in (('top', sup), ('start', izq), ('bottom', inf), ('end', der)):
        e = OxmlElement(f'w:{k}')
        e.set(qn('w:w'), str(v))
        e.set(qn('w:type'), 'dxa')
        m.append(e)
    tcPr.append(m)


def no_partir_fila(fila):
    trPr = fila._tr.get_or_add_trPr()
    e = OxmlElement('w:cantSplit')
    trPr.append(e)


def repetir_encabezado(fila):
    trPr = fila._tr.get_or_add_trPr()
    e = OxmlElement('w:tblHeader')
    trPr.append(e)


def campo(parrafo, instruccion, marcador='1'):
    run = parrafo.add_run()
    for tipo, texto in (('begin', None), ('instr', instruccion), ('separate', None), ('text', marcador), ('end', None)):
        if tipo == 'instr':
            e = OxmlElement('w:instrText')
            e.set(qn('xml:space'), 'preserve')
            e.text = f' {texto} '
        elif tipo == 'text':
            e = OxmlElement('w:t')
            e.text = texto
        else:
            e = OxmlElement('w:fldChar')
            e.set(qn('w:fldCharType'), tipo)
        run._r.append(e)
    return run


def borde_inferior(parrafo, color='0F7C86', tam=12):
    pPr = parrafo._p.get_or_add_pPr()
    pb = OxmlElement('w:pBdr')
    e = OxmlElement('w:bottom')
    e.set(qn('w:val'), 'single')
    e.set(qn('w:sz'), str(tam))
    e.set(qn('w:space'), '4')
    e.set(qn('w:color'), color)
    pb.append(e)
    pPr.append(pb)


def run(parrafo, texto, negrita=False, color=None, tam=None, cursiva=False):
    r = parrafo.add_run(texto)
    r.bold = negrita
    r.italic = cursiva
    if color is not None:
        r.font.color.rgb = color
    if tam:
        r.font.size = Pt(tam)
    return r


def texto_enriquecido(parrafo, texto, color=None, tam=None):
    """Admite **negrita** sencilla dentro del texto."""
    partes = texto.split('**')
    for i, parte in enumerate(partes):
        if parte:
            run(parrafo, parte, negrita=(i % 2 == 1), color=color, tam=tam)


# --------------------------------------------------------------------------- estilos
def configurar_estilos(doc):
    st = doc.styles
    normal = st['Normal']
    normal.font.name = FUENTE
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), FUENTE)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXTO
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    def titulo(nombre, tam, color, antes, despues, mayus=False):
        s = st[nombre]
        s.font.name = FUENTE
        rf = s.element.rPr.rFonts
        for att in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(att), FUENTE)
        for att in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            if rf.get(qn(att)) is not None:
                del rf.attrib[qn(att)]
        s.font.size = Pt(tam)
        s.font.bold = True
        s.font.italic = False
        s.font.color.rgb = color
        s.font.all_caps = mayus
        s.paragraph_format.space_before = Pt(antes)
        s.paragraph_format.space_after = Pt(despues)
        s.paragraph_format.keep_with_next = True

    titulo('Heading 1', 22, PRIMARIO_OSC, 0, 10)
    titulo('Heading 2', 14.5, PRIMARIO_OSC, 16, 6)
    titulo('Heading 3', 11.5, PRIMARIO, 12, 4)
    cap = st['Caption']
    cap.font.name = FUENTE
    cap.font.size = Pt(9)
    cap.font.italic = False
    cap.font.bold = False
    cap.font.color.rgb = TENUE
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(10)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for n in ('TOC 1', 'TOC 2'):
        try:
            s = st[n]
        except KeyError:
            from docx.enum.style import WD_STYLE_TYPE
            s = st.add_style(n, WD_STYLE_TYPE.PARAGRAPH)
        s.font.name = FUENTE
        s.font.size = Pt(11 if n == 'TOC 1' else 10)
        s.font.bold = n == 'TOC 1'
        s.paragraph_format.space_before = Pt(8 if n == 'TOC 1' else 0)
        s.paragraph_format.space_after = Pt(2)
        s.paragraph_format.left_indent = Cm(0 if n == 'TOC 1' else 0.6)


def encabezado_pie(seccion, titulo):
    seccion.header.is_linked_to_previous = False
    seccion.footer.is_linked_to_previous = False
    h = seccion.header.paragraphs[0]
    h.text = ''
    run(h, 'EcoLact', negrita=True, color=PRIMARIO, tam=8.5)
    run(h, f'   ·   {titulo}', color=TENUE, tam=8.5)
    borde_inferior(h, 'D5DDDF', 6)
    f = seccion.footer.paragraphs[0]
    f.text = ''
    f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(f, 'Página ', color=TENUE, tam=8.5)
    campo(f, 'PAGE').font.size = Pt(8.5)
    run(f, ' de ', color=TENUE, tam=8.5)
    campo(f, 'NUMPAGES').font.size = Pt(8.5)


# --------------------------------------------------------------------------- bloques
def tabla_ficha(doc, filas, ancho_etq=Cm(4.0)):
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    bordes_tabla(t, 'D5DDDF', 4)
    for etq, valor in filas:
        fila = t.add_row()
        no_partir_fila(fila)
        a, b = fila.cells
        a.width, b.width = ancho_etq, ANCHO_UTIL - ancho_etq
        sombrear(a, HEX_PRIMARIO_SUAVE)
        for c in (a, b):
            margenes_celda(c)
        pa = a.paragraphs[0]
        run(pa, etq.upper(), negrita=True, color=PRIMARIO_OSC, tam=8.5)
        valores = valor if isinstance(valor, list) else [valor]
        for i, v in enumerate(valores):
            p = b.paragraphs[0] if i == 0 else b.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            texto_enriquecido(p, ('• ' if isinstance(valor, list) else '') + v, tam=10)
    mantener_unida(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def tabla_datos(doc, encabezados, filas, anchos):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    bordes_tabla(t, 'D5DDDF', 4)
    cab = t.rows[0]
    repetir_encabezado(cab)
    for i, h in enumerate(encabezados):
        c = cab.cells[i]
        c.width = anchos[i]
        sombrear(c, '0B5961')
        margenes_celda(c, 70, 110, 70, 110)
        run(c.paragraphs[0], h, negrita=True, color=RGBColor(255, 255, 255), tam=9)
    for n, f in enumerate(filas):
        fila = t.add_row()
        no_partir_fila(fila)
        for i, v in enumerate(f):
            c = fila.cells[i]
            c.width = anchos[i]
            margenes_celda(c, 60, 110, 60, 110)
            if n % 2 == 1:
                sombrear(c, 'F4F7F7')
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            texto_enriquecido(p, str(v), tam=9)
    mantener_unida(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def mantener_unida(tabla, max_filas=14):
    """Evita encabezados huérfanos: tablas cortas se mantienen en una sola página."""
    filas = tabla.rows
    tope = len(filas) - 1 if len(filas) <= max_filas else 1
    for fila in filas[:tope]:
        for celda in fila.cells:
            for par in celda.paragraphs:
                par.paragraph_format.keep_with_next = True


def recuadro(doc, etiqueta, texto, color_fondo, color_etq):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    c.width = ANCHO_UTIL
    sombrear(c, color_fondo)
    margenes_celda(c, 90, 160, 90, 160)
    no_partir_fila(t.rows[0])
    tblPr = t._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    e = OxmlElement('w:left')
    e.set(qn('w:val'), 'single')
    e.set(qn('w:sz'), '24')
    e.set(qn('w:color'), color_etq)
    b.append(e)
    tblPr.append(b)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run(p, etiqueta + '  ', negrita=True, color=RGBColor.from_string(color_etq), tam=9.5)
    texto_enriquecido(p, texto, tam=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def lista(doc, items, vineta='•'):
    for it in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        p.paragraph_format.space_after = Pt(2)
        run(p, f'{vineta}  ', color=PRIMARIO, negrita=True)
        texto_enriquecido(p, it)


def copia_ligera(ruta_img):
    """Copia JPEG 4:4:4 (calidad 90) solo para incrustar en el Word: mantiene el texto nítido y
    reduce el tamaño del DOCX para que Word lo abra y exporte sin colgarse. El PNG queda como evidencia."""
    ruta_img = Path(ruta_img)
    destino = BASE / 'evidencias' / 'docx' / (ruta_img.parent.name + '_' + ruta_img.stem + '.jpg')
    if not destino.exists() or destino.stat().st_mtime < ruta_img.stat().st_mtime:
        destino.parent.mkdir(parents=True, exist_ok=True)
        with Image.open(ruta_img) as im:
            im.convert('RGB').save(destino, 'JPEG', quality=90, subsampling=0, optimize=True)
    return destino


def figura(doc, ruta_img, pie, ancho_max=ANCHO_UTIL):
    with Image.open(ruta_img) as im:
        w, h = im.size
    # Nunca ampliar más allá de ~110 ppp para no pixelar; nunca deformar (solo ancho).
    ancho = min(ancho_max, Emu(int(w / 110 * 914400)))
    alto_max = Cm(17.5)
    if ancho * h / w > alto_max:
        ancho = Emu(int(alto_max * w / h))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(copia_ligera(ruta_img)), width=ancho)
    c = doc.add_paragraph(style='Caption')
    run(c, 'Figura ', negrita=True, color=TENUE, tam=9)
    r = campo(c, 'SEQ Figura \\* ARABIC')
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = TENUE
    run(c, f'. {pie}', color=TENUE, tam=9)


# --------------------------------------------------------------------------- evidencias
_registros = {}
_uso = []  # trazabilidad captura -> procedimiento para evidencias/ledger.json


def captura(nombre_registro, nombre):
    if nombre_registro not in _registros:
        ruta = BASE / nombre_registro
        _registros[nombre_registro] = {c['nombre']: c for c in json.loads(ruta.read_text(encoding='utf-8'))['capturas']}
    try:
        return _registros[nombre_registro][nombre]
    except KeyError:
        raise SystemExit(f'Captura "{nombre}" no encontrada en {nombre_registro}')


def imagen_paso(cap_json, proc, n, paso, sufijo=None):
    cap = captura(paso.get('registro', proc.get('registro', cap_json['registro'])), paso['captura'])
    src = BASE / cap['archivo']
    dst = ANOTADAS / cap_json['id'] / f"{proc['codigo']}_{sufijo or f'P{n:02d}'}.png"
    objetivo = paso.get('objetivo')
    if objetivo:
        rect = list(cap['objetivos'][objetivo])
        # Mensajes a todo lo ancho: el recuadro se limita a la franja donde está su texto.
        if paso.get('limitar_x'):
            rect[0] = max(rect[0], paso['limitar_x'][0])
            rect[2] = min(rect[2], paso['limitar_x'][1])
        with Image.open(src) as im:
            W, H = im.size
        if rect[0] < 0 or rect[1] < 0 or rect[2] > W or rect[3] > H:
            raise SystemExit(f"{proc['codigo']} paso {n}: el control '{objetivo}' {rect} queda fuera de la captura {cap['nombre']} ({W}x{H}).")
        anotar(src, dst, rect, n, recorte=paso.get('recorte', 'auto'), etiqueta=paso.get('etiqueta'), posicion=paso.get('posicion'))
        _uso.append({'procedimiento': proc['codigo'], 'paso': n, 'captura': cap['nombre'], 'original': cap['archivo'],
                     'anotada': dst.relative_to(BASE).as_posix(), 'objetivo': objetivo, 'rectangulo': rect,
                     'url': cap.get('url'), 'estado': proc['estado']})
    else:
        dst.parent.mkdir(parents=True, exist_ok=True)
        with Image.open(src) as im:
            caja = paso.get('recorte')
            (im.crop(tuple(caja)) if isinstance(caja, list) else im).convert('RGB').save(dst, optimize=True)
    return dst


# --------------------------------------------------------------------------- partes del manual
def portada(doc, m):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    banda = doc.add_table(rows=1, cols=1)
    c = banda.rows[0].cells[0]
    c.width = ANCHO_UTIL
    sombrear(c, '0B5961')
    margenes_celda(c, 200, 300, 200, 300)
    p = c.paragraphs[0]
    run(p, 'MANUAL DE USUARIO', negrita=True, color=RGBColor(0xE1, 0xF2, 0xF1), tam=11)
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(BASE / 'recursos' / 'vaquita.png'), width=Cm(5.2))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    run(p, 'EcoLact', negrita=True, color=PRIMARIO_OSC, tam=40)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, m['subtitulo'], color=ACENTO, tam=17, negrita=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    run(p, m['descripcion'], color=TENUE, tam=11)
    for _ in range(6):
        doc.add_paragraph()
    tabla_ficha(doc, [
        ('Sistema', m['sistema']),
        ('Rol documentado', 'Administrador'),
        ('Versión del manual', m['version']),
        ('Fecha', m['fecha']),
        ('Estado', m['estado']),
    ], ancho_etq=Cm(4.5))


def titulo_sin_indice(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run(p, texto, negrita=True, color=PRIMARIO_OSC, tam=16)
    borde_inferior(p)


def control_e_indice(doc, m):
    titulo_sin_indice(doc, 'Control del documento')
    tabla_datos(doc, ['Versión', 'Fecha', 'Descripción del cambio'],
                [[v['version'], v['fecha'], v['cambio']] for v in m['historial']],
                [Cm(2.4), Cm(2.8), Cm(10.8)])
    titulo_sin_indice(doc, 'Contenido')
    p = doc.add_paragraph()
    campo(p, 'TOC \\o "1-2" \\h \\z \\u', 'Haga clic derecho y elija «Actualizar campo» para generar el índice.')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def capitulo_general(doc, cap, num):
    h = doc.add_paragraph(style='Heading 1')
    run(h, f'{num}. {cap["titulo"]}')
    borde_inferior(h)
    for bloque in cap['bloques']:
        tipo = bloque['tipo']
        if tipo == 'h2':
            doc.add_paragraph(bloque['texto'], style='Heading 2')
        elif tipo == 'h3':
            doc.add_paragraph(bloque['texto'], style='Heading 3')
        elif tipo == 'p':
            texto_enriquecido(doc.add_paragraph(), bloque['texto'])
        elif tipo == 'lista':
            lista(doc, bloque['items'])
        elif tipo == 'tabla':
            anchos = [Cm(a) for a in bloque['anchos']]
            tabla_datos(doc, bloque['encabezados'], bloque['filas'], anchos)
        elif tipo == 'nota':
            recuadro(doc, bloque.get('etiqueta', 'NOTA'), bloque['texto'], HEX_AVISO_SUAVE, 'E6A23C')
        elif tipo == 'figura':
            cap_ = captura(bloque['registro'], bloque['captura'])
            figura(doc, BASE / cap_['archivo'], bloque['pie'])
        elif tipo == 'salto':
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def procedimiento(doc, cap, proc, num_cap, num_proc):
    h = doc.add_paragraph(style='Heading 2')
    run(h, f'{num_cap}.{num_proc}  {proc["titulo"]}')
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    run(p, proc['codigo'], negrita=True, color=ACENTO, tam=9)
    run(p, '   ·   ' + proc['estado'], color=TENUE, tam=9)
    tabla_ficha(doc, [
        ('Objetivo', proc['objetivo']),
        ('Acceso', proc['acceso']),
        ('Requisitos', proc['requisitos']),
    ])

    for n, paso in enumerate(proc['pasos'], start=1):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        run(p, f'PASO {n:02d}', negrita=True, color=RGBColor(0xD9, 0x26, 0x26), tam=10)
        run(p, '   ' + paso['titulo'], negrita=True, color=PRIMARIO_OSC, tam=11)
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        texto_enriquecido(p, paso['instruccion'])
        if paso.get('captura'):
            figura(doc, imagen_paso(cap, proc, n, paso), paso['pie'])
        if paso.get('resultado'):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            run(p, '✓ Resultado: ', negrita=True, color=RGBColor(0x27, 0x9B, 0x75), tam=10)
            texto_enriquecido(p, paso['resultado'], tam=10)
        if paso.get('nota'):
            recuadro(doc, 'IMPORTANTE', paso['nota'], HEX_AVISO_SUAVE, 'E6A23C')

    doc.add_paragraph('Resultado final', style='Heading 3')
    recuadro(doc, 'RESULTADO', proc['resultado'], HEX_EXITO_SUAVE, '279B75')
    if proc.get('validaciones'):
        doc.add_paragraph('Validaciones y mensajes del sistema', style='Heading 3')
        tabla_datos(doc, ['Campo / condición', 'Regla', 'Mensaje mostrado'], proc['validaciones'],
                    [Cm(3.8), Cm(5.4), Cm(6.8)])
    for k, fig in enumerate(proc.get('figuras_validacion', []), start=1):
        figura(doc, imagen_paso(cap, proc, '!', fig, sufijo=f'V{k:02d}'), fig['pie'])
    if proc.get('errores'):
        doc.add_paragraph('Errores frecuentes y solución', style='Heading 3')
        tabla_datos(doc, ['Situación', 'Cómo resolverla'], proc['errores'], [Cm(6.5), Cm(9.5)])


def capitulo_modulo(doc, cap, num):
    h = doc.add_paragraph(style='Heading 1')
    run(h, f'{num}. {cap["titulo"]}')
    borde_inferior(h)
    texto_enriquecido(doc.add_paragraph(), cap['introduccion'])
    if cap.get('ficha'):
        tabla_ficha(doc, [(k, v) for k, v in cap['ficha']])
    if cap.get('pantalla'):
        doc.add_paragraph('Elementos de la pantalla', style='Heading 3')
        pant = cap['pantalla']
        cap_ = captura(pant.get('registro', cap['registro']), pant['captura'])
        figura(doc, BASE / cap_['archivo'], pant['pie'])
        tabla_datos(doc, ['Zona de la pantalla', 'Función'], pant['elementos'], [Cm(4.6), Cm(11.4)])
    for i, proc in enumerate(cap['procedimientos'], start=1):
        procedimiento(doc, cap, proc, num, i)


ESTADO_COLOR = {'Verificado en interfaz': 'E8F8EF', 'Verificado parcialmente': 'FFF4DF'}


def anexo_verificacion(doc, cap, capitulos, letra):
    h = doc.add_paragraph(style='Heading 1')
    run(h, f'Anexo {letra}. {cap["titulo"]}')
    borde_inferior(h)
    texto_enriquecido(doc.add_paragraph(), cap['introduccion'])
    doc.add_paragraph('Procedimientos documentados', style='Heading 2')
    filas = [[p['codigo'], p['titulo'], c['titulo'], p['estado']] for c in capitulos for p in c.get('procedimientos', [])]
    t = tabla_datos(doc, ['Código', 'Procedimiento', 'Capítulo', 'Estado'], filas, [Cm(2.6), Cm(5.6), Cm(3.6), Cm(4.2)])
    for fila in t.rows[1:]:
        estado = fila.cells[3].text
        sombrear(fila.cells[3], ESTADO_COLOR.get(estado, 'FDECEC'))
    for bloque in cap.get('secciones', []):
        doc.add_paragraph(bloque['titulo'], style='Heading 2')
        if bloque.get('texto'):
            texto_enriquecido(doc.add_paragraph(), bloque['texto'])
        if bloque.get('tabla'):
            tb = bloque['tabla']
            tabla_datos(doc, tb['encabezados'], tb['filas'], [Cm(a) for a in tb['anchos']])
        if bloque.get('items'):
            lista(doc, bloque['items'])


def construir(sin_pdf=False):
    m = json.loads((MANUAL / 'manual.json').read_text(encoding='utf-8'))
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2.0)
    sec.header_distance = sec.footer_distance = Cm(1.0)
    configurar_estilos(doc)
    doc.core_properties.title = m['titulo']
    doc.core_properties.subject = m['subtitulo']
    doc.core_properties.author = 'EcoLact — Documentación'
    doc.core_properties.language = 'es-PE'

    portada(doc, m)
    nueva = doc.add_section(WD_SECTION.NEW_PAGE)
    nueva.different_first_page_header_footer = False
    encabezado_pie(nueva, m['subtitulo'])
    # La sección 0 (portada) no lleva encabezado ni pie: su primera página usa el encabezado vacío.
    control_e_indice(doc, m)

    capitulos = []
    for num, ruta in enumerate(m['capitulos'], start=1):
        cap = json.loads((MANUAL / ruta).read_text(encoding='utf-8'))
        if num > 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        (capitulo_modulo if cap.get('tipo') == 'modulo' else capitulo_general)(doc, cap, num)
        capitulos.append(cap)
    for letra, ruta in zip('ABCDEF', m.get('anexos', [])):
        cap = json.loads((MANUAL / ruta).read_text(encoding='utf-8'))
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        anexo_verificacion(doc, cap, capitulos, letra)

    (BASE / 'evidencias' / 'ledger.json').write_text(json.dumps({
        'documento': 'Registro de evidencias de EcoLact — Manual del Administrador',
        'criterio': 'Cada imagen anotada procede de una captura original medida en el DOM (capturar.py). Sin datos personales: nombres, DNI y usuario difuminados.',
        'generado': date.today().isoformat(),
        'evidencias': _uso,
    }, ensure_ascii=False, indent=2), encoding='utf-8')
    SALIDA.mkdir(exist_ok=True)
    docx = SALIDA / f'{NOMBRE}.docx'
    doc.save(docx)
    print('DOCX:', docx)
    if not sin_pdf:
        ps = BASE / 'scripts' / 'word_a_pdf.ps1'
        subprocess.run(['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(ps), str(docx)], check=True, timeout=1500)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--sin-pdf', action='store_true')
    construir(ap.parse_args().sin_pdf)
