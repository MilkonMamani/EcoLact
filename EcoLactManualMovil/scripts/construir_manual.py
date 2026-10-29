#!/usr/bin/env python3
"""Construye el Manual de Usuario de EcoLact Móvil (Word editable + PDF) desde manual/contenido.py.

Estilo visual de uso rápido: cada procedimiento es título + una frase + una pantalla por hoja: captura a la
izquierda y sus pasos numerados a la derecha (+ una nota breve). Hojas blancas, texto negro, sin fondos ni tablas en los
procedimientos; códigos y estados de verificación solo en el Anexo D. Las capturas conservan sus colores;
las marcas son rojas.

Uso: python scripts/construir_manual.py [--sin-pdf]
"""
import argparse
import json
import re
import shutil
import subprocess
import sys
import unicodedata
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE / 'scripts'))
sys.path.insert(0, str(BASE / 'manual'))
from anotar import anotar  # noqa: E402
import contenido as C  # noqa: E402

PROTEGIDAS = BASE / 'trabajo' / 'protegidas'
ORIGINALES = BASE / 'evidencias' / 'originales'
ANOTADAS = BASE / 'evidencias' / 'anotadas'
DOCX_IMG = BASE / 'trabajo' / 'docx_img'
SALIDA = BASE / 'salida'
NOMBRE = 'Manual_Usuario_EcoLact_Movil'
OCR = json.loads((BASE / 'trabajo' / 'ocr_palabras.json').read_text(encoding='utf-8'))

NEGRO = RGBColor(0, 0, 0)
GRIS = RGBColor(0x40, 0x40, 0x40)
HEX_BORDE = '000000'
HEX_GRIS_CLARO = 'F2F2F2'
FUENTE = 'Arial'
ANCHO = Cm(16.6)
ANCHO_FIG = Cm(7.0)      # una captura por hoja
HEX_GRIS_LINEA = 'A6A6A6'
SECCION = {}             # código MOB-… → número de sección (p. ej. 3.5)

ROL_DE_CAPITULO = {'GEN': 'general', 'PLA': 'planta', 'VEN': 'ventas', 'CAL': 'calidad', 'PRV': 'proveedor', 'ACO': 'acopiador'}


# --------------------------------------------------------------------------- XML auxiliares
def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def bordes_tabla(tabla, tam=4, internos=True, color=HEX_BORDE):
    tblPr = tabla._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for lado in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement(f'w:{lado}')
        visible = internos or lado not in ('insideH', 'insideV')
        e.set(qn('w:val'), 'single' if visible else 'nil')
        if visible:
            e.set(qn('w:sz'), str(tam))
            e.set(qn('w:color'), color)
        b.append(e)
    tblPr.append(b)


def sin_bordes(tabla):
    bordes_tabla(tabla, internos=False)
    for e in tabla._tbl.tblPr.find(qn('w:tblBorders')):
        e.set(qn('w:val'), 'nil')


def margenes_celda(celda, sup=70, izq=110, inf=70, der=110):
    tcPr = celda._tc.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for k, v in (('top', sup), ('start', izq), ('bottom', inf), ('end', der)):
        e = OxmlElement(f'w:{k}')
        e.set(qn('w:w'), str(v))
        e.set(qn('w:type'), 'dxa')
        m.append(e)
    tcPr.append(m)


def fila_fija(fila, encabezado=False):
    trPr = fila._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))
    if encabezado:
        trPr.append(OxmlElement('w:tblHeader'))


def anchos_columnas(tabla, anchos):
    tabla.autofit = False
    tblPr = tabla._tbl.tblPr
    lay = OxmlElement('w:tblLayout')
    lay.set(qn('w:type'), 'fixed')
    tblPr.append(lay)
    for fila in tabla.rows:
        for c, a in zip(fila.cells, anchos):
            c.width = a


def campo(parrafo, instruccion, marcador='1'):
    r = parrafo.add_run()
    for tipo, texto in (('begin', None), ('instr', instruccion), ('separate', None), ('text', marcador), ('end', None)):
        if tipo == 'instr':
            e = OxmlElement('w:instrText')
            e.set(qn('xml:space'), 'preserve')
            e.text = f' {texto} '
        elif tipo == 'text':
            e = OxmlElement('w:t')
            e.text = texto
        else:
            e = OxmlElement('w:fldChar')
            e.set(qn('w:fldCharType'), tipo)
        r._r.append(e)
    return r


def borde_parrafo(parrafo, lado='bottom', tam=8, color='000000', espacio=4):
    pPr = parrafo._p.get_or_add_pPr()
    pb = pPr.find(qn('w:pBdr'))
    if pb is None:
        pb = OxmlElement('w:pBdr')
        pPr.append(pb)
    e = OxmlElement(f'w:{lado}')
    e.set(qn('w:val'), 'single')
    e.set(qn('w:sz'), str(tam))
    e.set(qn('w:space'), str(espacio))
    e.set(qn('w:color'), color)
    pb.append(e)


def run(p, texto, negrita=False, tam=None, cursiva=False, color=NEGRO):
    r = p.add_run(texto)
    r.bold, r.italic = negrita, cursiva
    r.font.color.rgb = color
    if tam:
        r.font.size = Pt(tam)
    return r


def secciones(texto):
    return re.sub(r'\{(MOB-[A-Z]+-\d+)\}', lambda m: SECCION[m.group(1)], texto)


def rico(p, texto, tam=None, color=NEGRO):
    """Admite **negrita** y referencias {MOB-…} dentro del texto."""
    texto = secciones(texto)
    for i, parte in enumerate(texto.split('**')):
        if parte:
            run(p, parte, negrita=i % 2 == 1, tam=tam, color=color)


def parrafo(doc_o_celda, texto='', tam=None, antes=0, despues=4, alinear=None, con_siguiente=False):
    p = doc_o_celda.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(antes), Pt(despues)
    pf.keep_with_next = con_siguiente
    if alinear:
        p.alignment = alinear
    if texto:
        rico(p, texto, tam)
    return p


def primer_parrafo(celda, despues=0):
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(despues)
    return p


# --------------------------------------------------------------------------- estilos
def configurar_estilos(doc):
    st = doc.styles
    n = st['Normal']
    n.font.name = FUENTE
    n.element.rPr.rFonts.set(qn('w:eastAsia'), FUENTE)
    n.font.size = Pt(10.5)
    n.font.color.rgb = NEGRO
    n.paragraph_format.space_after = Pt(4)
    n.paragraph_format.line_spacing = 1.1

    def titulo(nombre, tam, antes, despues):
        s = st[nombre]
        s.font.name = FUENTE
        rf = s.element.rPr.rFonts
        for att in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(att), FUENTE)
        for att in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            if rf.get(qn(att)) is not None:
                del rf.attrib[qn(att)]
        s.font.size = Pt(tam)
        s.font.bold = True
        s.font.italic = False
        s.font.color.rgb = NEGRO
        s.paragraph_format.space_before = Pt(antes)
        s.paragraph_format.space_after = Pt(despues)
        s.paragraph_format.keep_with_next = True

    titulo('Heading 1', 17, 0, 10)
    titulo('Heading 2', 12.5, 10, 2)
    titulo('Heading 3', 11, 10, 4)
    cap = st['Caption']
    cap.font.name = FUENTE
    cap.font.size = Pt(8.5)
    cap.font.italic = False
    cap.font.bold = False
    cap.font.color.rgb = GRIS
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(2)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    from docx.enum.style import WD_STYLE_TYPE
    for nombre, tam, neg, sang, antes in (('TOC 1', 10.5, True, 0, 6), ('TOC 2', 10, False, 0.6, 0)):
        try:
            s = st[nombre]
        except KeyError:
            s = st.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
        s.font.name = FUENTE
        s.font.size = Pt(tam)
        s.font.bold = neg
        s.font.color.rgb = NEGRO
        s.paragraph_format.left_indent = Cm(sang)
        s.paragraph_format.space_before = Pt(antes)
        s.paragraph_format.space_after = Pt(1)


def encabezado_pie(seccion):
    seccion.header.is_linked_to_previous = False
    seccion.footer.is_linked_to_previous = False
    h = seccion.header.paragraphs[0]
    h.text = ''
    run(h, 'EcoLact', negrita=True, tam=8.5)
    run(h, '   ·   Manual de Usuario de la aplicación móvil', tam=8.5, color=GRIS)
    borde_parrafo(h, tam=4)
    f = seccion.footer.paragraphs[0]
    f.text = ''
    borde_parrafo(f, 'top', tam=4)
    f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(f, 'Ecolácteos Huata · Versión 3.0 · Octubre de 2026      ', tam=8.5, color=GRIS)
    run(f, 'Página ', tam=8.5, color=GRIS)
    campo(f, 'PAGE').font.size = Pt(8.5)
    run(f, ' de ', tam=8.5, color=GRIS)
    campo(f, 'NUMPAGES').font.size = Pt(8.5)


# --------------------------------------------------------------------------- bloques de texto
def tabla(doc, encabezados, filas, anchos_cm, tam=9.5):
    t = doc.add_table(rows=1 if encabezados else 0, cols=len(anchos_cm))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    bordes_tabla(t)
    if encabezados:
        cab = t.rows[0]
        fila_fija(cab, encabezado=True)
        for c, h in zip(cab.cells, encabezados):
            sombrear(c, HEX_GRIS_CLARO)
            margenes_celda(c)
            p = primer_parrafo(c)
            p.paragraph_format.keep_with_next = True
            run(p, h, negrita=True, tam=tam)
    for f in filas:
        fila = t.add_row()
        fila_fija(fila)
        for c, v in zip(fila.cells, f):
            margenes_celda(c)
            valores = v if isinstance(v, list) else [v]
            for i, val in enumerate(valores):
                p = primer_parrafo(c) if i == 0 else c.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                rico(p, ('• ' if isinstance(v, list) else '') + str(val), tam)
    anchos_columnas(t, [Cm(a) for a in anchos_cm])
    if len(t.rows) <= 12:
        for fila in t.rows[:-1]:
            for c in fila.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True
    parrafo(doc, despues=2)
    return t


def ficha(doc, filas, ancho_etq=3.6):
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    bordes_tabla(t)
    for etq, valor in filas:
        fila = t.add_row()
        fila_fija(fila)
        a, b = fila.cells
        sombrear(a, HEX_GRIS_CLARO)
        for c in (a, b):
            margenes_celda(c)
        run(primer_parrafo(a), etq, negrita=True, tam=9.5)
        valores = valor if isinstance(valor, list) else [valor]
        for i, v in enumerate(valores):
            p = primer_parrafo(b) if i == 0 else b.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            rico(p, ('• ' if isinstance(valor, list) else '') + v, 9.5)
        for c in (a, b):
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True
    anchos_columnas(t, [Cm(ancho_etq), ANCHO - Cm(ancho_etq)])
    # La ficha queda unida a lo que sigue (la captura del procedimiento), nunca sola al pie de página.
    parrafo(doc, despues=2, con_siguiente=True)


def recuadro(doc, etiqueta, texto, detalle=None):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    bordes_tabla(t, tam=6)
    c = t.rows[0].cells[0]
    margenes_celda(c, 90, 140, 90, 140)
    fila_fija(t.rows[0])
    p = primer_parrafo(c)
    run(p, etiqueta.upper() + '. ', negrita=True, tam=9.5)
    rico(p, texto, 10)
    if detalle:
        p2 = c.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.space_before = Pt(3)
        run(p2, 'Evidencia: ', negrita=True, tam=8.5, color=GRIS)
        run(p2, detalle, tam=8.5, color=GRIS, cursiva=True)
    anchos_columnas(t, [ANCHO])
    parrafo(doc, despues=2)


def lista(doc, items):
    for it in items:
        p = parrafo(doc, despues=2)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        run(p, '•  ', negrita=True)
        rico(p, it)


def titulo_sin_indice(doc, texto):
    p = parrafo(doc, antes=6, despues=8, con_siguiente=True)
    run(p, texto, negrita=True, tam=15)
    borde_parrafo(p, tam=8)


def h1(doc, texto):
    h = doc.add_paragraph(style='Heading 1')
    run(h, texto)
    borde_parrafo(h, tam=8)


# --------------------------------------------------------------------------- capturas
def _norm(t):
    t = unicodedata.normalize('NFKD', t).encode('ascii', 'ignore').decode().lower()
    return ''.join(ch for ch in t if ch.isalnum() or ch == '-' or ch == ':')


def ubicar(img, frase):
    """Rectángulo (x1,y1,x2,y2) de una frase en la captura, según el OCR guardado."""
    ocur = 1
    if '#' in frase:
        frase, ocur = frase.split('#')
        ocur = int(ocur)
    objetivo = [_norm(w) for w in frase.split()]
    palabras = [(_norm(w[0]), w[1:]) for w in OCR[img + '.png']]
    hallados = []
    for i in range(len(palabras)):
        ok = True
        for k, obj in enumerate(objetivo):
            if i + k >= len(palabras):
                ok = False
                break
            w = palabras[i + k][0]
            if not (w.startswith(obj) or (len(obj) > 3 and obj in w)):
                ok = False
                break
        if ok:
            cajas = [palabras[i + k][1] for k in range(len(objetivo))]
            hallados.append((min(c[0] for c in cajas), min(c[1] for c in cajas),
                             max(c[2] for c in cajas), max(c[3] for c in cajas)))
    if len(hallados) < ocur:
        raise SystemExit(f'No se encontró «{frase}» (aparición {ocur}) en {img}. Ajuste el ancla en contenido.py.')
    return hallados[ocur - 1]


def botones_verdes(img):
    """Franjas de botones verdes sólidos (de arriba abajo): el OCR no lee su texto blanco."""
    import numpy as np
    a = np.asarray(Image.open(PROTEGIDAS / f'{img}.png').convert('RGB')).astype(int)
    r, g, b = a[..., 0], a[..., 1], a[..., 2]
    m = (g > 100) & (g - r > 50) & (g - b > 20) & (r < 90)
    filas = m.sum(1) > a.shape[1] * 0.35
    out, y = [], 0
    while y < len(filas):
        if filas[y]:
            y0 = y
            while y < len(filas) and filas[y]:
                y += 1
            if y - y0 > 30:
                xs = np.where(m[y0:y].any(0))[0]
                out.append((int(xs.min()), y0, int(xs.max()), y))
        y += 1
    return out


def rect_marca(img, ancla, ancho_img):
    tipo = ancla[0]
    k = ancho_img / 912
    x_izq, x_der = round(30 * k), round(ancho_img - 30 * k)
    if tipo == 'caja':
        return tuple(ancla[1:5])
    if tipo == 'boton':
        bandas = botones_verdes(img)
        if len(bandas) < ancla[1]:
            raise SystemExit(f'{img}: no hay botón verde n.° {ancla[1]}')
        x1, y1, x2, y2 = bandas[ancla[1] - 1]
        return (x1 - 6, y1 - 6, x2 + 6, y2 + 6)
    if tipo == 'fila':
        x1, y1, x2, y2 = ubicar(img, ancla[1])
        alto = ancla[2] if len(ancla) > 2 else 45
        cy = (y1 + y2) // 2
        return (x_izq, cy - alto, x_der, cy + alto)
    if tipo == 'texto':
        x1, y1, x2, y2 = ubicar(img, ancla[1])
        m = ancla[2] if len(ancla) > 2 else 16
        return (x1 - m, y1 - m, x2 + m, y2 + m)
    if tipo == 'zona':
        a = ubicar(img, ancla[1])
        b = ubicar(img, ancla[2])
        m = ancla[3] if len(ancla) > 3 else 26
        return (x_izq, min(a[1], b[1]) - m, x_der, max(a[3], b[3]) + m)
    raise SystemExit(f'Ancla desconocida: {ancla}')


_uso = []


def preparar_imagen(fig, proc, rol):
    img = fig['img']
    origen = PROTEGIDAS / f'{img}.png'
    (ORIGINALES / rol).mkdir(parents=True, exist_ok=True)
    shutil.copyfile(origen, ORIGINALES / rol / f'{img}.png')
    (ANOTADAS / rol).mkdir(parents=True, exist_ok=True)
    destino = ANOTADAS / rol / f"{proc['codigo']}_{img}.png"
    if fig.get('marcas'):
        with Image.open(origen) as im:
            w = im.width
        marcas = [(n, rect_marca(img, a, w)) for n, a in fig['marcas']]
        anotar(origen, destino, marcas)
    else:
        marcas = []
        shutil.copyfile(origen, destino)
    _uso.append({'procedimiento': proc['codigo'], 'captura': img, 'original': f'evidencias/originales/{rol}/{img}.png',
                 'anotada': destino.relative_to(BASE).as_posix(),
                 'marcas': [[n, list(r)] for n, r in marcas] if marcas else 'incrustadas en la captura original',
                 'fuente': 'Manual_Movil_Actual.pdf (captura de la aplicación del 28/10/2026)'})
    DOCX_IMG.mkdir(parents=True, exist_ok=True)
    jpg = DOCX_IMG / (destino.stem + '.jpg')
    with Image.open(destino) as im:
        im = im.convert('RGB')
        from PIL import ImageDraw
        ImageDraw.Draw(im).rectangle([0, 0, im.width - 1, im.height - 1], outline=(150, 150, 150), width=2)
        im.save(jpg, 'JPEG', quality=90, subsampling=0, optimize=True)
    return jpg


def celda_figura(celda, ruta, ancho):
    p = primer_parrafo(celda)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ruta), width=ancho)


def paso(contenedor, n, texto, primero=False, tam=10.5):
    """Un paso: número en negrita con sangría francesa y una instrucción breve."""
    p = contenedor.paragraphs[0] if primero else contenedor.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(0), Pt(7)
    pf.left_indent, pf.first_line_indent = Cm(0.75), Cm(-0.75)
    pf.tab_stops.add_tab_stop(Cm(0.75))
    pf.keep_together = True
    run(p, f'{n}\t', negrita=True, tam=tam)
    rico(p, texto, tam)
    return p


def captura_pendiente(contenedor):
    p = contenedor.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Cm(0.75)
    run(p, 'Captura pendiente.', cursiva=True, tam=8.5, color=GRIS)
    return p


def pantallas(bloque):
    """Divide un bloque en pantallas: cada captura con los pasos cuyos números lleva marcados."""
    figs, pasos = bloque.get('figs', []), bloque['pasos']
    if not figs:
        return [(None, pasos)]
    grupos = []
    for f in figs:
        nums = f.get('nums') or [n for n, _ in (f.get('marcas') or [])]
        grupos.append((f, [s for s in pasos if s[0] in nums]))
    sueltos = [s for s in pasos if all(s not in g for _, g in grupos)]
    if sueltos:
        f, g = grupos[-1]
        grupos[-1] = (f, sorted(g + sueltos))
    return grupos


def pantalla(doc, fig, pasos, proc, rol):
    """Una captura grande a la izquierda y sus pasos a la derecha; sin captura, lista con aviso."""
    if fig is None:
        # Pasos sin evidencia visual: lista simple y aviso discreto (detalle en el Anexo D).
        for n, texto in pasos:
            pf = paso(doc, n, texto).paragraph_format
            pf.keep_with_next, pf.space_after = True, Pt(4)
        captura_pendiente(doc)
        return
    ruta = preparar_imagen(fig, proc, rol)
    anchos = [ANCHO_FIG + Cm(0.6), ANCHO - ANCHO_FIG - Cm(0.6)]
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    sin_bordes(t)
    fila_fija(t.rows[0])
    imagen, texto = t.rows[0].cells
    margenes_celda(imagen, 40, 40, 40, 40)
    imagen.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    celda_figura(imagen, ruta, ANCHO_FIG)
    margenes_celda(texto, 60, 260, 40, 40)
    texto.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for k, (n, txt) in enumerate(pasos):
        paso(texto, n, txt, primero=k == 0, tam=11)
    anchos_columnas(t, anchos)
    sep = parrafo(doc, despues=2)
    sep.paragraph_format.line_spacing = Pt(4)   # separador mínimo tras la captura


# --------------------------------------------------------------------------- estructura
def portada(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    p = parrafo(doc, antes=10, despues=2)
    run(p, 'ECOLÁCTEOS HUATA', negrita=True, tam=12)
    p = parrafo(doc, despues=90)
    run(p, 'Sistema EcoLact', tam=10.5, color=GRIS)
    borde_parrafo(p, tam=12, espacio=8)
    p = parrafo(doc, despues=4)
    run(p, 'Manual de Usuario', negrita=True, tam=34)
    p = parrafo(doc, despues=40)
    run(p, 'Aplicación móvil EcoLact', tam=20)
    for i, rol in enumerate(['Encargado de planta', 'Personal de ventas', 'Analista de calidad', 'Proveedor', 'Acopiador'], 3):
        p = parrafo(doc, despues=3)
        run(p, f'{i}   ', negrita=True, tam=11)
        run(p, rol, tam=11)
    for _ in range(14):
        parrafo(doc, despues=10)
    p = parrafo(doc, despues=2)
    borde_parrafo(p, 'top', tam=4, espacio=6)
    run(p, 'Versión 3.0  ·  29 de octubre de 2026', tam=10)
    p = parrafo(doc, despues=0)
    run(p, 'Manual de uso rápido por roles · Android', tam=10, color=GRIS)


def control_e_indice(doc):
    titulo_sin_indice(doc, 'Control del documento')
    tabla(doc, ['Versión', 'Fecha', 'Descripción del cambio'], [
        ['1.0', '10/2026', 'Manual anterior: cinco guías visuales por rol.'],
        ['2.0', '29/10/2026', 'Revisión contra el código actual; procedimientos codificados; anexos de verificación.'],
        ['3.0', '29/10/2026', 'Rediseño visual de uso rápido: captura y pasos numerados por procedimiento; '
                              'códigos y estados de verificación trasladados al Anexo D.'],
    ], [2.0, 2.6, 12.0])
    titulo_sin_indice(doc, 'Índice')
    p = doc.add_paragraph()
    campo(p, 'TOC \\o "1-2" \\h \\z \\u', 'Haga clic derecho y elija «Actualizar campo» para generar el índice.')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def capitulo_introduccion(doc, num):
    h1(doc, f'{num}. {C.INTRODUCCION["titulo"]}')
    for tipo, *resto in C.INTRODUCCION['bloques']:
        if tipo == 'h2':
            doc.add_paragraph(f'{num}.{_sub(num)}  {resto[0]}', style='Heading 2')
        elif tipo == 'p':
            parrafo(doc, resto[0], despues=6)
        elif tipo == 'lista':
            lista(doc, resto[0])


_contador = {}


def _sub(num):
    _contador[num] = _contador.get(num, 0) + 1
    return _contador[num]


def subtitulo(doc, texto):
    p = parrafo(doc, antes=6, despues=3, con_siguiente=True)
    run(p, texto, negrita=True, tam=10.5)


def nota(doc, texto):
    p = parrafo(doc, antes=2, despues=4)
    p.paragraph_format.keep_together = True
    p.paragraph_format.left_indent = Cm(0.25)
    borde_parrafo(p, 'left', tam=12, color='000000', espacio=6)
    if texto.startswith('Importante:'):
        run(p, 'Importante:', negrita=True, tam=10)
        rico(p, texto[len('Importante:'):], 10)
    else:
        rico(p, texto, 10)


def procedimiento(doc, proc, num_cap, num_proc, rol):
    h = doc.add_paragraph(f'{num_cap}.{num_proc}  {proc["titulo"]}', style='Heading 2')
    # Un procedimiento por hoja; el primero comparte la hoja con la presentación del capítulo.
    h.paragraph_format.page_break_before = num_proc > 1
    borde_parrafo(h, tam=4, color=HEX_GRIS_LINEA, espacio=2)
    parrafo(doc, proc['frase'], antes=3, despues=4, con_siguiente=True)
    # Una pantalla por hoja: cada captura adicional abre hoja con el título del procedimiento como referencia.
    hay_captura = False
    for bloque in proc['bloques']:
        for fig, pasos in pantallas(bloque):
            if fig is not None and hay_captura:
                c = parrafo(doc, despues=6, con_siguiente=True)
                c.paragraph_format.page_break_before = True
                run(c, f'{num_cap}.{num_proc}  {proc["titulo"]} (continuación)', negrita=True, tam=10, color=GRIS)
                borde_parrafo(c, tam=4, color=HEX_GRIS_LINEA, espacio=2)
            pantalla(doc, fig, pasos, proc, rol)
            hay_captura = hay_captura or fig is not None
    if proc.get('lista'):
        titulo, items = proc['lista']
        subtitulo(doc, titulo)
        lista(doc, items)
        for item in doc.paragraphs[-len(items):]:   # la lista no se parte entre páginas
            item.paragraph_format.keep_with_next = True
    if proc.get('nota'):
        nota(doc, proc['nota'])


def capitulo_rol(doc, cap, num):
    h1(doc, f'{num}. {cap["titulo"]}')
    parrafo(doc, cap['intro'], despues=6)
    for etq, valor in cap.get('datos', []):
        p = parrafo(doc, despues=2)
        run(p, etq + ': ', negrita=True)
        rico(p, valor)
    if cap.get('flujo'):
        subtitulo(doc, 'Tu trabajo del día')
        for i, t in enumerate(cap['flujo'], 1):
            paso(doc, i, t).paragraph_format.space_after = Pt(2)
    if cap.get('stock'):
        subtitulo(doc, 'Cómo cambia el stock')
        lista(doc, cap['stock'])
    rol = ROL_DE_CAPITULO[cap['procedimientos'][0]['codigo'].split('-')[1]]
    for i, proc in enumerate(cap['procedimientos'], 1):
        procedimiento(doc, proc, num, i, rol)


def anexos(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    h1(doc, 'Anexo A. Glosario')
    tabla(doc, ['Término', 'Definición'], C.GLOSARIO, [4.4, 12.2])

    h1(doc, 'Anexo B. Problemas frecuentes')
    tabla(doc, ['Situación', 'Qué hacer'], [[a, secciones(b)] for a, b in C.PROBLEMAS], [7.0, 9.6])

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    h1(doc, 'Anexo C. Mensajes de la aplicación')
    parrafo(doc, 'Mensajes que muestra la aplicación cuando falta un dato o no se puede completar una operación.', despues=6)
    filas = [[SECCION[p['codigo']], sit, msj] for cap in C.CAPITULOS for p in cap['procedimientos']
             for sit, msj in p.get('mensajes', [])]
    tabla(doc, ['Sección', 'Situación', 'Mensaje'], filas, [1.6, 5.0, 10.0], tam=9)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    h1(doc, 'Anexo D. Seguimiento técnico')
    parrafo(doc, 'Código y estado de verificación de cada procedimiento. Las capturas proceden de la aplicación en '
                 'ejecución del 28/10/2026 y se contrastaron con el código fuente del 29/10/2026.', despues=6)
    procs = [p for cap in C.CAPITULOS for p in cap['procedimientos']]
    tabla(doc, ['Sección', 'Código', 'Procedimiento', 'Estado'],
          [[SECCION[p['codigo']], p['codigo'], p['titulo'], p['estado']] for p in procs], [1.6, 2.7, 6.4, 5.9], tam=8.5)
    cuenta = {}
    for p in procs:
        cuenta[p['estado']] = cuenta.get(p['estado'], 0) + 1
    tabla(doc, ['Estado', 'Procedimientos'], [[k, str(v)] for k, v in cuenta.items()] + [['**Total**', f'**{len(procs)}**']],
          [12.0, 4.6], tam=9)
    subtitulo(doc, 'Capturas pendientes')
    tabla(doc, ['Sección', 'Código', 'Pantalla o prueba que falta'],
          [[SECCION[p['codigo']], p['codigo'], p['pendiente']] for p in procs if p.get('pendiente')], [1.6, 2.7, 12.3], tam=9)
    subtitulo(doc, 'Limitaciones')
    lista(doc, C.LIMITACIONES)


def construir(sin_pdf=False):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
    sec.header_distance = sec.footer_distance = Cm(1.0)
    configurar_estilos(doc)
    cp = doc.core_properties
    cp.title = 'Manual de Usuario — Aplicación móvil EcoLact'
    cp.subject = 'Manual por roles: Encargado de planta, Personal de ventas, Analista de calidad, Proveedor y Acopiador'
    cp.author = 'EcoLact — Documentación'
    cp.language = 'es-PE'

    portada(doc)
    nueva = doc.add_section(WD_SECTION.NEW_PAGE)
    nueva.different_first_page_header_footer = False
    encabezado_pie(nueva)
    control_e_indice(doc)

    for num, cap in enumerate(C.CAPITULOS, start=2):
        for k, proc in enumerate(cap['procedimientos'], 1):
            SECCION[proc['codigo']] = f'{num}.{k}'
    capitulo_introduccion(doc, 1)
    for num, cap in enumerate(C.CAPITULOS, start=2):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        capitulo_rol(doc, cap, num)
    anexos(doc)

    (BASE / 'evidencias' / 'ledger.json').write_text(json.dumps({
        'documento': 'Registro de evidencias — Manual de Usuario EcoLact Móvil',
        'criterio': 'Capturas auténticas de la aplicación (28/10/2026) extraídas del manual anterior; DNI, teléfonos, usuarios y QR '
                    'pixelados (scripts/privacidad.py). Marcas ubicadas sobre el texto real de cada control (OCR).',
        'generado': date.today().isoformat(),
        'evidencias': _uso,
    }, ensure_ascii=False, indent=2), encoding='utf-8')
    SALIDA.mkdir(exist_ok=True)
    docx = SALIDA / f'{NOMBRE}.docx'
    doc.save(docx)
    print('DOCX:', docx)
    if not sin_pdf:
        ps = BASE / 'scripts' / 'word_a_pdf.ps1'
        subprocess.run(['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(ps), str(docx)], check=True, timeout=1500)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--sin-pdf', action='store_true')
    construir(ap.parse_args().sin_pdf)
