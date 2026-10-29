#!/usr/bin/env python3
"""Construye el Manual de Usuario de EcoLact Móvil (Word editable + PDF) desde manual/contenido.py.

Estilo editorial: hojas blancas, texto y títulos negros, tablas y recuadros con bordes finos negros,
sin fondos de color. Las capturas conservan los colores de la aplicación; solo las marcas son rojas.

Uso: python scripts/construir_manual.py [--sin-pdf]
"""
import argparse
import json
import shutil
import subprocess
import sys
import unicodedata
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE / 'scripts'))
sys.path.insert(0, str(BASE / 'manual'))
from anotar import anotar  # noqa: E402
import contenido as C  # noqa: E402

PROTEGIDAS = BASE / 'trabajo' / 'protegidas'
ORIGINALES = BASE / 'evidencias' / 'originales'
ANOTADAS = BASE / 'evidencias' / 'anotadas'
DOCX_IMG = BASE / 'trabajo' / 'docx_img'
SALIDA = BASE / 'salida'
NOMBRE = 'Manual_Usuario_EcoLact_Movil'
OCR = json.loads((BASE / 'trabajo' / 'ocr_palabras.json').read_text(encoding='utf-8'))

NEGRO = RGBColor(0, 0, 0)
GRIS = RGBColor(0x40, 0x40, 0x40)
HEX_BORDE = '000000'
HEX_GRIS_CLARO = 'F2F2F2'
FUENTE = 'Arial'
ANCHO = Cm(16.6)
ANCHO_FIG = Cm(5.8)

ROL_DE_CAPITULO = {'GEN': 'general', 'PLA': 'planta', 'VEN': 'ventas', 'CAL': 'calidad', 'PRV': 'proveedor', 'ACO': 'acopiador'}


# --------------------------------------------------------------------------- XML auxiliares
def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def bordes_tabla(tabla, tam=4, internos=True, color=HEX_BORDE):
    tblPr = tabla._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for lado in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement(f'w:{lado}')
        visible = internos or lado not in ('insideH', 'insideV')
        e.set(qn('w:val'), 'single' if visible else 'nil')
        if visible:
            e.set(qn('w:sz'), str(tam))
            e.set(qn('w:color'), color)
        b.append(e)
    tblPr.append(b)


def sin_bordes(tabla):
    bordes_tabla(tabla, internos=False)
    for e in tabla._tbl.tblPr.find(qn('w:tblBorders')):
        e.set(qn('w:val'), 'nil')


def margenes_celda(celda, sup=70, izq=110, inf=70, der=110):
    tcPr = celda._tc.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for k, v in (('top', sup), ('start', izq), ('bottom', inf), ('end', der)):
        e = OxmlElement(f'w:{k}')
        e.set(qn('w:w'), str(v))
        e.set(qn('w:type'), 'dxa')
        m.append(e)
    tcPr.append(m)


def fila_fija(fila, encabezado=False):
    trPr = fila._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))
    if encabezado:
        trPr.append(OxmlElement('w:tblHeader'))


def anchos_columnas(tabla, anchos):
    tabla.autofit = False
    tblPr = tabla._tbl.tblPr
    lay = OxmlElement('w:tblLayout')
    lay.set(qn('w:type'), 'fixed')
    tblPr.append(lay)
    for fila in tabla.rows:
        for c, a in zip(fila.cells, anchos):
            c.width = a


def campo(parrafo, instruccion, marcador='1'):
    r = parrafo.add_run()
    for tipo, texto in (('begin', None), ('instr', instruccion), ('separate', None), ('text', marcador), ('end', None)):
        if tipo == 'instr':
            e = OxmlElement('w:instrText')
            e.set(qn('xml:space'), 'preserve')
            e.text = f' {texto} '
        elif tipo == 'text':
            e = OxmlElement('w:t')
            e.text = texto
        else:
            e = OxmlElement('w:fldChar')
            e.set(qn('w:fldCharType'), tipo)
        r._r.append(e)
    return r


def borde_parrafo(parrafo, lado='bottom', tam=8, color='000000', espacio=4):
    pPr = parrafo._p.get_or_add_pPr()
    pb = pPr.find(qn('w:pBdr'))
    if pb is None:
        pb = OxmlElement('w:pBdr')
        pPr.append(pb)
    e = OxmlElement(f'w:{lado}')
    e.set(qn('w:val'), 'single')
    e.set(qn('w:sz'), str(tam))
    e.set(qn('w:space'), str(espacio))
    e.set(qn('w:color'), color)
    pb.append(e)


def run(p, texto, negrita=False, tam=None, cursiva=False, color=NEGRO):
    r = p.add_run(texto)
    r.bold, r.italic = negrita, cursiva
    r.font.color.rgb = color
    if tam:
        r.font.size = Pt(tam)
    return r


def rico(p, texto, tam=None, color=NEGRO):
    """Admite **negrita** dentro del texto."""
    for i, parte in enumerate(texto.split('**')):
        if parte:
            run(p, parte, negrita=i % 2 == 1, tam=tam, color=color)


def parrafo(doc_o_celda, texto='', tam=None, antes=0, despues=4, alinear=None, con_siguiente=False):
    p = doc_o_celda.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(antes), Pt(despues)
    pf.keep_with_next = con_siguiente
    if alinear:
        p.alignment = alinear
    if texto:
        rico(p, texto, tam)
    return p


def primer_parrafo(celda, despues=0):
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(despues)
    return p


# --------------------------------------------------------------------------- estilos
def configurar_estilos(doc):
    st = doc.styles
    n = st['Normal']
    n.font.name = FUENTE
    n.element.rPr.rFonts.set(qn('w:eastAsia'), FUENTE)
    n.font.size = Pt(10.5)
    n.font.color.rgb = NEGRO
    n.paragraph_format.space_after = Pt(4)
    n.paragraph_format.line_spacing = 1.1

    def titulo(nombre, tam, antes, despues):
        s = st[nombre]
        s.font.name = FUENTE
        rf = s.element.rPr.rFonts
        for att in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(att), FUENTE)
        for att in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            if rf.get(qn(att)) is not None:
                del rf.attrib[qn(att)]
        s.font.size = Pt(tam)
        s.font.bold = True
        s.font.italic = False
        s.font.color.rgb = NEGRO
        s.paragraph_format.space_before = Pt(antes)
        s.paragraph_format.space_after = Pt(despues)
        s.paragraph_format.keep_with_next = True

    titulo('Heading 1', 17, 0, 10)
    titulo('Heading 2', 13, 14, 5)
    titulo('Heading 3', 11, 10, 4)
    cap = st['Caption']
    cap.font.name = FUENTE
    cap.font.size = Pt(8.5)
    cap.font.italic = False
    cap.font.bold = False
    cap.font.color.rgb = GRIS
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(2)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    from docx.enum.style import WD_STYLE_TYPE
    for nombre, tam, neg, sang, antes in (('TOC 1', 10.5, True, 0, 6), ('TOC 2', 10, False, 0.6, 0)):
        try:
            s = st[nombre]
        except KeyError:
            s = st.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
        s.font.name = FUENTE
        s.font.size = Pt(tam)
        s.font.bold = neg
        s.font.color.rgb = NEGRO
        s.paragraph_format.left_indent = Cm(sang)
        s.paragraph_format.space_before = Pt(antes)
        s.paragraph_format.space_after = Pt(1)


def encabezado_pie(seccion):
    seccion.header.is_linked_to_previous = False
    seccion.footer.is_linked_to_previous = False
    h = seccion.header.paragraphs[0]
    h.text = ''
    run(h, 'EcoLact', negrita=True, tam=8.5)
    run(h, '   ·   Manual de Usuario de la Aplicación Móvil', tam=8.5, color=GRIS)
    borde_parrafo(h, tam=4)
    f = seccion.footer.paragraphs[0]
    f.text = ''
    borde_parrafo(f, 'top', tam=4)
    f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(f, 'Ecolácteos Huata · Versión 2.0 · Octubre de 2026      ', tam=8.5, color=GRIS)
    run(f, 'Página ', tam=8.5, color=GRIS)
    campo(f, 'PAGE').font.size = Pt(8.5)
    run(f, ' de ', tam=8.5, color=GRIS)
    campo(f, 'NUMPAGES').font.size = Pt(8.5)


# --------------------------------------------------------------------------- bloques de texto
def tabla(doc, encabezados, filas, anchos_cm, tam=9.5):
    t = doc.add_table(rows=1 if encabezados else 0, cols=len(anchos_cm))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    bordes_tabla(t)
    if encabezados:
        cab = t.rows[0]
        fila_fija(cab, encabezado=True)
        for c, h in zip(cab.cells, encabezados):
            sombrear(c, HEX_GRIS_CLARO)
            margenes_celda(c)
            p = primer_parrafo(c)
            p.paragraph_format.keep_with_next = True
            run(p, h, negrita=True, tam=tam)
    for f in filas:
        fila = t.add_row()
        fila_fija(fila)
        for c, v in zip(fila.cells, f):
            margenes_celda(c)
            valores = v if isinstance(v, list) else [v]
            for i, val in enumerate(valores):
                p = primer_parrafo(c) if i == 0 else c.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                rico(p, ('• ' if isinstance(v, list) else '') + str(val), tam)
    anchos_columnas(t, [Cm(a) for a in anchos_cm])
    if len(t.rows) <= 12:
        for fila in t.rows[:-1]:
            for c in fila.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True
    parrafo(doc, despues=2)
    return t


def ficha(doc, filas, ancho_etq=3.6):
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    bordes_tabla(t)
    for etq, valor in filas:
        fila = t.add_row()
        fila_fija(fila)
        a, b = fila.cells
        sombrear(a, HEX_GRIS_CLARO)
        for c in (a, b):
            margenes_celda(c)
        run(primer_parrafo(a), etq, negrita=True, tam=9.5)
        valores = valor if isinstance(valor, list) else [valor]
        for i, v in enumerate(valores):
            p = primer_parrafo(b) if i == 0 else b.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            rico(p, ('• ' if isinstance(valor, list) else '') + v, 9.5)
        for c in (a, b):
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True
    anchos_columnas(t, [Cm(ancho_etq), ANCHO - Cm(ancho_etq)])
    # La ficha queda unida a lo que sigue (la captura del procedimiento), nunca sola al pie de página.
    parrafo(doc, despues=2, con_siguiente=True)


def recuadro(doc, etiqueta, texto, detalle=None):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    bordes_tabla(t, tam=6)
    c = t.rows[0].cells[0]
    margenes_celda(c, 90, 140, 90, 140)
    fila_fija(t.rows[0])
    p = primer_parrafo(c)
    run(p, etiqueta.upper() + '. ', negrita=True, tam=9.5)
    rico(p, texto, 10)
    if detalle:
        p2 = c.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.space_before = Pt(3)
        run(p2, 'Evidencia: ', negrita=True, tam=8.5, color=GRIS)
        run(p2, detalle, tam=8.5, color=GRIS, cursiva=True)
    anchos_columnas(t, [ANCHO])
    parrafo(doc, despues=2)


def lista(doc, items):
    for it in items:
        p = parrafo(doc, despues=2)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        run(p, '•  ', negrita=True)
        rico(p, it)


def titulo_sin_indice(doc, texto):
    p = parrafo(doc, antes=6, despues=8, con_siguiente=True)
    run(p, texto, negrita=True, tam=15)
    borde_parrafo(p, tam=8)


def h1(doc, texto):
    h = doc.add_paragraph(style='Heading 1')
    run(h, texto)
    borde_parrafo(h, tam=8)


# --------------------------------------------------------------------------- capturas
def _norm(t):
    t = unicodedata.normalize('NFKD', t).encode('ascii', 'ignore').decode().lower()
    return ''.join(ch for ch in t if ch.isalnum() or ch == '-' or ch == ':')


def ubicar(img, frase):
    """Rectángulo (x1,y1,x2,y2) de una frase en la captura, según el OCR guardado."""
    ocur = 1
    if '#' in frase:
        frase, ocur = frase.split('#')
        ocur = int(ocur)
    objetivo = [_norm(w) for w in frase.split()]
    palabras = [(_norm(w[0]), w[1:]) for w in OCR[img + '.png']]
    hallados = []
    for i in range(len(palabras)):
        ok = True
        for k, obj in enumerate(objetivo):
            if i + k >= len(palabras):
                ok = False
                break
            w = palabras[i + k][0]
            if not (w.startswith(obj) or (len(obj) > 3 and obj in w)):
                ok = False
                break
        if ok:
            cajas = [palabras[i + k][1] for k in range(len(objetivo))]
            hallados.append((min(c[0] for c in cajas), min(c[1] for c in cajas),
                             max(c[2] for c in cajas), max(c[3] for c in cajas)))
    if len(hallados) < ocur:
        raise SystemExit(f'No se encontró «{frase}» (aparición {ocur}) en {img}. Ajuste el ancla en contenido.py.')
    return hallados[ocur - 1]


def botones_verdes(img):
    """Franjas de botones verdes sólidos (de arriba abajo): el OCR no lee su texto blanco."""
    import numpy as np
    a = np.asarray(Image.open(PROTEGIDAS / f'{img}.png').convert('RGB')).astype(int)
    r, g, b = a[..., 0], a[..., 1], a[..., 2]
    m = (g > 100) & (g - r > 50) & (g - b > 20) & (r < 90)
    filas = m.sum(1) > a.shape[1] * 0.35
    out, y = [], 0
    while y < len(filas):
        if filas[y]:
            y0 = y
            while y < len(filas) and filas[y]:
                y += 1
            if y - y0 > 30:
                xs = np.where(m[y0:y].any(0))[0]
                out.append((int(xs.min()), y0, int(xs.max()), y))
        y += 1
    return out


def rect_marca(img, ancla, ancho_img):
    tipo = ancla[0]
    k = ancho_img / 912
    x_izq, x_der = round(30 * k), round(ancho_img - 30 * k)
    if tipo == 'caja':
        return tuple(ancla[1:5])
    if tipo == 'boton':
        bandas = botones_verdes(img)
        if len(bandas) < ancla[1]:
            raise SystemExit(f'{img}: no hay botón verde n.° {ancla[1]}')
        x1, y1, x2, y2 = bandas[ancla[1] - 1]
        return (x1 - 6, y1 - 6, x2 + 6, y2 + 6)
    if tipo == 'fila':
        x1, y1, x2, y2 = ubicar(img, ancla[1])
        alto = ancla[2] if len(ancla) > 2 else 45
        cy = (y1 + y2) // 2
        return (x_izq, cy - alto, x_der, cy + alto)
    if tipo == 'texto':
        x1, y1, x2, y2 = ubicar(img, ancla[1])
        m = ancla[2] if len(ancla) > 2 else 16
        return (x1 - m, y1 - m, x2 + m, y2 + m)
    if tipo == 'zona':
        a = ubicar(img, ancla[1])
        b = ubicar(img, ancla[2])
        m = ancla[3] if len(ancla) > 3 else 26
        return (x_izq, min(a[1], b[1]) - m, x_der, max(a[3], b[3]) + m)
    raise SystemExit(f'Ancla desconocida: {ancla}')


_uso = []


def preparar_imagen(fig, proc, rol):
    img = fig['img']
    origen = PROTEGIDAS / f'{img}.png'
    (ORIGINALES / rol).mkdir(parents=True, exist_ok=True)
    shutil.copyfile(origen, ORIGINALES / rol / f'{img}.png')
    (ANOTADAS / rol).mkdir(parents=True, exist_ok=True)
    destino = ANOTADAS / rol / f"{proc['codigo']}_{img}.png"
    if fig.get('marcas'):
        with Image.open(origen) as im:
            w = im.width
        marcas = [(n, rect_marca(img, a, w)) for n, a in fig['marcas']]
        anotar(origen, destino, marcas)
    else:
        marcas = []
        shutil.copyfile(origen, destino)
    _uso.append({'procedimiento': proc['codigo'], 'captura': img, 'original': f'evidencias/originales/{rol}/{img}.png',
                 'anotada': destino.relative_to(BASE).as_posix(),
                 'marcas': [[n, list(r)] for n, r in marcas] if marcas else 'incrustadas en la captura original',
                 'fuente': 'Manual_Movil_Actual.pdf (captura de la aplicación del 28/10/2026)'})
    DOCX_IMG.mkdir(parents=True, exist_ok=True)
    jpg = DOCX_IMG / (destino.stem + '.jpg')
    with Image.open(destino) as im:
        im.convert('RGB').save(jpg, 'JPEG', quality=90, subsampling=0, optimize=True)
    return jpg


def celda_figura(celda, ruta, pie, ancho=ANCHO_FIG):
    p = primer_parrafo(celda)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ruta), width=ancho)
    c = celda.add_paragraph(style='Caption')
    run(c, 'Figura ', negrita=True, tam=8.5, color=GRIS)
    r = campo(c, 'SEQ Figura \\* ARABIC')
    r.bold = True
    r.font.size = Pt(8.5)
    run(c, f'. {pie}', tam=8.5, color=GRIS)


def lista_pasos(contenedor, pasos, es_elemento, tam=9.5):
    """Pasos o elementos con su número de marca, dentro de una celda o del documento."""
    for i, (n, titulo, texto) in enumerate(pasos):
        if hasattr(contenedor, 'paragraphs') and i == 0 and not contenedor.paragraphs[0].text and len(contenedor.paragraphs) == 1:
            p = contenedor.paragraphs[0]
        else:
            p = contenedor.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        sang = 0.7 if es_elemento else 1.75
        p.paragraph_format.left_indent = Cm(sang)
        p.paragraph_format.first_line_indent = Cm(-sang)
        p.paragraph_format.keep_together = True
        etiqueta = f'{n}' if es_elemento else f'Paso {n:02d}'
        run(p, etiqueta + '\t', negrita=True, tam=tam)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(sang))
        run(p, titulo + '. ', negrita=True, tam=tam)
        rico(p, texto, tam)


def bloque_figuras(doc, bloque, proc, rol):
    figs = bloque.get('figs', [])
    pasos = bloque.get('pasos') or bloque.get('elementos') or []
    es_elem = 'elementos' in bloque
    if not figs:
        t = doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        bordes_tabla(t, tam=4)
        c = t.rows[0].cells[0]
        margenes_celda(c, 90, 140, 40, 140)
        lista_pasos(c, pasos, es_elem)
        anchos_columnas(t, [ANCHO])
        parrafo(doc, despues=2)
        return
    rutas = [preparar_imagen(f, proc, rol) for f in figs]
    if len(figs) == 1:
        t = doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        bordes_tabla(t, tam=4)
        a, b = t.rows[0].cells
        fila_fija(t.rows[0])
        margenes_celda(a, 100, 100, 80, 100)
        margenes_celda(b, 120, 160, 80, 140)
        celda_figura(a, rutas[0], figs[0]['pie'])
        if es_elem:
            run(primer_parrafo(b, 6), 'Elementos señalados', negrita=True, tam=10)
        lista_pasos(b, pasos, es_elem)
        b.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        anchos_columnas(t, [Cm(6.5), ANCHO - Cm(6.5)])
    else:
        t = doc.add_table(rows=2, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        bordes_tabla(t, tam=4)
        for k in range(2):
            c = t.rows[0].cells[k]
            margenes_celda(c, 100, 100, 80, 100)
            celda_figura(c, rutas[k], figs[k]['pie'])
        fila_fija(t.rows[0])
        abajo = t.rows[1].cells[0].merge(t.rows[1].cells[1])
        margenes_celda(abajo, 120, 160, 60, 160)
        if es_elem:
            run(primer_parrafo(abajo, 6), 'Elementos señalados', negrita=True, tam=10)
        lista_pasos(abajo, pasos, es_elem)
        anchos_columnas(t, [ANCHO / 2, ANCHO / 2])
    parrafo(doc, despues=2)


# --------------------------------------------------------------------------- estructura
def portada(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    p = parrafo(doc, antes=10, despues=2)
    run(p, 'ECOLÁCTEOS HUATA', negrita=True, tam=12)
    p = parrafo(doc, despues=60)
    run(p, 'Documentación del sistema EcoLact', tam=10.5, color=GRIS)
    borde_parrafo(p, tam=12, espacio=8)
    p = parrafo(doc, despues=4)
    run(p, 'Manual de Usuario', negrita=True, tam=34)
    p = parrafo(doc, despues=26)
    run(p, 'Aplicación móvil EcoLact', tam=20)
    p = parrafo(doc, despues=6)
    run(p, 'Roles documentados', negrita=True, tam=11)
    for i, rol in enumerate(['Encargado de planta', 'Personal de ventas', 'Analista de calidad', 'Proveedor', 'Acopiador'], 3):
        p = parrafo(doc, despues=2)
        run(p, f'Capítulo {i}   ', negrita=True, tam=10.5)
        run(p, rol, tam=10.5)
    for _ in range(6):
        parrafo(doc, despues=10)
    ficha(doc, [
        ('Sistema', 'EcoLact — aplicación móvil Android (Kotlin Multiplatform)'),
        ('Tipo de documento', 'Manual de usuario por roles'),
        ('Versión', '2.0'),
        ('Fecha', '29 de octubre de 2026'),
        ('Estado', 'Revisado contra el código actual de la aplicación; ver Anexo C.'),
    ], ancho_etq=4.2)


def control_e_indice(doc):
    titulo_sin_indice(doc, 'Control del documento')
    tabla(doc, ['Versión', 'Fecha', 'Descripción del cambio'], [
        ['1.0', '10/2026', 'Manual anterior: cinco guías visuales por rol (60 páginas).'],
        ['2.0', '29/10/2026', 'Revisión contra el código actual; corrección de textos, capturas y marcas; procedimientos codificados; '
                              'funciones nuevas (recorrido, muestreo, análisis, liquidaciones, boleta, sincronización); diseño formal en blanco y negro; anexos de verificación.'],
    ], [2.0, 2.6, 12.0])
    titulo_sin_indice(doc, 'Índice general')
    p = doc.add_paragraph()
    campo(p, 'TOC \\o "1-2" \\h \\z \\u', 'Haga clic derecho y elija «Actualizar campo» para generar el índice.')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def capitulo_introduccion(doc, num):
    h1(doc, f'{num}. {C.INTRODUCCION["titulo"]}')
    for tipo, *resto in C.INTRODUCCION['bloques']:
        if tipo == 'h2':
            doc.add_paragraph(f'{num}.{_sub(num)}  {resto[0]}', style='Heading 2')
        elif tipo == 'p':
            parrafo(doc, resto[0])
        elif tipo == 'lista':
            lista(doc, resto[0])
        elif tipo == 'tabla':
            tabla(doc, resto[0], resto[1], resto[2])
        elif tipo == 'nota':
            recuadro(doc, resto[0], resto[1])


_contador = {}


def _sub(num):
    _contador[num] = _contador.get(num, 0) + 1
    return _contador[num]


def procedimiento(doc, proc, num_cap, num_proc, rol):
    doc.add_paragraph(f'{num_cap}.{num_proc}  {proc["titulo"]}', style='Heading 2')
    p = parrafo(doc, despues=4, con_siguiente=True)
    run(p, 'Procedimiento ' + proc['codigo'], negrita=True, tam=9.5)
    run(p, '   ·   Estado: ' + proc['estado'], tam=9.5, color=GRIS)
    ficha(doc, [('Propósito', proc['proposito']), ('Acceso', proc['acceso']), ('Requisitos', proc['requisitos'])])
    for bloque in proc['bloques']:
        if 'tabla' in bloque:
            enc, filas, anchos = bloque['tabla']
            tabla(doc, enc, filas, anchos)
        else:
            bloque_figuras(doc, bloque, proc, rol)
    texto, evidencia = proc['resultado']
    recuadro(doc, 'Resultado esperado', texto, evidencia)
    if proc.get('validaciones'):
        p = parrafo(doc, antes=4, despues=3, con_siguiente=True)
        run(p, 'Validaciones y mensajes del sistema', negrita=True, tam=10.5)
        tabla(doc, ['Condición', 'Mensaje o regla'], proc['validaciones'], [5.6, 11.0], tam=9)
    for nota in proc.get('notas', []):
        recuadro(doc, 'Observación', nota)


def capitulo_rol(doc, cap, num):
    h1(doc, f'{num}. {cap["titulo"]}')
    parrafo(doc, cap['intro'])
    if cap.get('ficha'):
        ficha(doc, cap['ficha'])
    if cap.get('flujo'):
        p = parrafo(doc, antes=4, despues=3, con_siguiente=True)
        run(p, 'Flujo de trabajo', negrita=True, tam=11)
        tabla(doc, ['N.°', 'Actividad', 'Procedimiento'], cap['flujo'], [1.2, 10.4, 5.0])
    if cap.get('stock'):
        p = parrafo(doc, antes=4, despues=3, con_siguiente=True)
        run(p, 'Cómo cambia el stock', negrita=True, tam=11)
        tabla(doc, ['Operación', 'Efecto en el inventario'], cap['stock'], [5.6, 11.0])
    rol = ROL_DE_CAPITULO[cap['procedimientos'][0]['codigo'].split('-')[1]]
    for i, proc in enumerate(cap['procedimientos'], 1):
        procedimiento(doc, proc, num, i, rol)


def anexos(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    h1(doc, 'Anexo A. Glosario')
    tabla(doc, ['Término', 'Definición'], C.GLOSARIO, [4.4, 12.2])

    h1(doc, 'Anexo B. Problemas frecuentes')
    tabla(doc, ['Situación', 'Qué hacer'], C.PROBLEMAS, [7.0, 9.6])

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    h1(doc, 'Anexo C. Estado de verificación')
    parrafo(doc, 'La tabla resume cómo se comprobó cada procedimiento (criterios en la sección 1.5). Las capturas proceden de '
                 'la aplicación en ejecución del 28/10/2026 y se contrastaron con el código fuente del 29/10/2026.')
    filas = [[p['codigo'], p['titulo'], p['estado']] for cap in C.CAPITULOS for p in cap['procedimientos']]
    tabla(doc, ['Código', 'Procedimiento', 'Estado'], filas, [2.8, 8.4, 5.4], tam=9)
    cuenta = {}
    for f in filas:
        cuenta[f[2]] = cuenta.get(f[2], 0) + 1
    tabla(doc, ['Estado', 'Procedimientos'], [[k, str(v)] for k, v in cuenta.items()] + [['Total', str(len(filas))]], [11.0, 5.6])

    h1(doc, 'Anexo D. Limitaciones')
    lista(doc, [
        'Durante la revisión solo se autorizó navegar: no se confirmaron ventas, compras, producciones, entregas, muestreos, análisis ni comunicados nuevos. '
        'Los resultados que no aparecen en las capturas se marcan como previstos.',
        'La cuenta de prueba disponible en el emulador tenía la sesión vencida y la base de datos local estaba casi vacía; por eso no se obtuvieron capturas nuevas '
        'y se reutilizaron las capturas auténticas del 28/10/2026, contrastadas con el código actual.',
        'Las capturas del Encargado de planta conservan las marcas numeradas de su edición original, porque no existen copias sin marcar.',
        'En algunas capturas se ve la barra flotante del teclado del emulador (Calidad: muestreo directo y nuevo comunicado). No se retocó para no alterar la interfaz.',
        'La lectura automática (OCR) de boletas y tickets depende del servidor y de la calidad de la foto; los valores deben revisarse siempre antes de confirmar.',
        'Al publicar un comunicado por «ZONA» o «INDIVIDUAL», la aplicación solicita identificadores numéricos en lugar de una lista para elegir.',
    ])
    h1(doc, 'Anexo E. Funcionalidades pendientes de verificación')
    tabla(doc, ['Rol', 'Funcionalidad', 'Qué falta comprobar'], [
        ['Encargado de planta', 'Escaneo de boleta (cámara, «Datos detectados»)', 'Capturas de la cámara y de la revisión OCR; confirmación con datos de prueba.'],
        ['Personal de ventas', 'Confirmación «Venta registrada» y boleta PDF', 'Capturas de la confirmación, la boleta, «Descargar PDF» y «Compartir».'],
        ['Personal de ventas', 'Registro de comprador; comunicados', 'Capturas del formulario «Nuevo comprador» y de la lista de comunicados.'],
        ['Analista de calidad', 'Muestreo aleatorio y selección manual', 'Capturas de «Muestreo», «Elegir muestra al azar» y «Seleccionar manualmente».'],
        ['Analista de calidad', 'Registro y confirmación del análisis', 'Capturas de «Tomar foto del ticket», «Datos detectados», «Resultado estimado» y «Análisis confirmado».'],
        ['Analista de calidad', 'Historial de análisis y proveedores', 'Capturas de ambas pantallas y su detalle.'],
        ['Proveedor', 'Detalle de pago (liquidación), detalle de compra, lista de comunicados, sincronización', 'Capturas con una liquidación generada y con datos de compras.'],
        ['Acopiador', 'Recorrido del día', 'Capturas de «Recorrido de hoy», paradas, «No entregó» y «Finalizar recorrido» con un recorrido asignado.'],
        ['Acopiador', 'Corrección de entregas con error', 'Captura de «Corregir entrega» ante un error de sincronización.'],
        ['Todos', 'Envío automático tras recuperar la conexión', 'Prueba funcional en un entorno de pruebas.'],
    ], [3.4, 6.2, 7.0], tam=9)


def construir(sin_pdf=False):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    sec.header_distance = sec.footer_distance = Cm(1.0)
    configurar_estilos(doc)
    cp = doc.core_properties
    cp.title = 'Manual de Usuario — Aplicación móvil EcoLact'
    cp.subject = 'Manual por roles: Encargado de planta, Personal de ventas, Analista de calidad, Proveedor y Acopiador'
    cp.author = 'EcoLact — Documentación'
    cp.language = 'es-PE'

    portada(doc)
    nueva = doc.add_section(WD_SECTION.NEW_PAGE)
    nueva.different_first_page_header_footer = False
    encabezado_pie(nueva)
    control_e_indice(doc)

    capitulo_introduccion(doc, 1)
    for num, cap in enumerate(C.CAPITULOS, start=2):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        capitulo_rol(doc, cap, num)
    anexos(doc)

    (BASE / 'evidencias' / 'ledger.json').write_text(json.dumps({
        'documento': 'Registro de evidencias — Manual de Usuario EcoLact Móvil',
        'criterio': 'Capturas auténticas de la aplicación (28/10/2026) extraídas del manual anterior; DNI, teléfonos, usuarios y QR '
                    'pixelados (scripts/privacidad.py). Marcas ubicadas sobre el texto real de cada control (OCR).',
        'generado': date.today().isoformat(),
        'evidencias': _uso,
    }, ensure_ascii=False, indent=2), encoding='utf-8')
    SALIDA.mkdir(exist_ok=True)
    docx = SALIDA / f'{NOMBRE}.docx'
    doc.save(docx)
    print('DOCX:', docx)
    if not sin_pdf:
        ps = BASE / 'scripts' / 'word_a_pdf.ps1'
        subprocess.run(['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(ps), str(docx)], check=True, timeout=1500)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--sin-pdf', action='store_true')
    construir(ap.parse_args().sin_pdf)
